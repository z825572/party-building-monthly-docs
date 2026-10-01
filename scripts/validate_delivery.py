#!/usr/bin/env python
"""Validate a monthly DOCX delivery folder before final handoff."""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from pathlib import Path


PLACEHOLDER_RE = re.compile(r"(?i)(?:x{2,}党支部|x{2,})")
MONTH_RE = re.compile(r"(\d{1,2})月")


def configure_stdout() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")


def import_dependencies():
    try:
        from docx import Document
    except ImportError as exc:
        raise SystemExit(
            "Missing dependency. Run this script with the bundled workspace "
            "Python that includes python-docx."
        ) from exc
    return Document


def branch_label(value: str | None) -> str:
    label = (value or "").strip()
    if not label:
        return ""
    if not label.endswith("党支部"):
        label += "党支部"
    return label


def find_docx(directory: Path) -> list[Path]:
    return sorted(
        path
        for path in directory.rglob("*.docx")
        if not path.name.startswith("~$")
        and ".tmp." not in path.name
        and path.is_file()
    )


def document_text(Document, path: Path) -> tuple[list[str], list[str]]:
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    styles = sorted(
        {
            paragraph.style.name
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        }
    )
    return paragraphs, styles


def validate_docx(Document, path: Path, month: int | None, branch: str) -> dict:
    errors = []
    warnings = []
    name = path.name

    if month is not None:
        months_in_name = [int(value) for value in MONTH_RE.findall(name)]
        if months_in_name and month not in months_in_name:
            errors.append(f"文件名月份与目标月份 {month}月 不一致。")
        elif not months_in_name:
            warnings.append(f"文件名没有显式写入月份 {month}月。")
    if branch and branch not in name:
        warnings.append(f"文件名未包含党支部名称 {branch}。")
    if PLACEHOLDER_RE.search(name):
        errors.append("文件名仍含 xxxx 占位符。")

    try:
        with zipfile.ZipFile(path) as archive:
            bad_member = archive.testzip()
    except (OSError, zipfile.BadZipFile) as exc:
        return {
            "path": str(path),
            "errors": [f"DOCX ZIP 无法读取: {exc}"],
            "warnings": warnings,
            "paragraphs": 0,
            "styles": [],
        }
    if bad_member:
        errors.append(f"DOCX ZIP 成员损坏: {bad_member}")

    try:
        paragraphs, styles = document_text(Document, path)
    except Exception as exc:
        errors.append(f"python-docx 无法重新打开文件: {exc}")
        paragraphs = []
        styles = []

    text = "\n".join(paragraphs).strip()
    if not text:
        errors.append("正文为空。")
    if PLACEHOLDER_RE.search(text):
        errors.append("正文仍含 xxxx 占位符。")

    return {
        "path": str(path),
        "errors": errors,
        "warnings": warnings,
        "paragraphs": len(paragraphs),
        "styles": styles,
    }


def load_expected_names(plan_path: Path) -> list[str]:
    try:
        plan = json.loads(plan_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise SystemExit(f"Cannot read plan JSON: {exc}") from exc
    activities = plan.get("activities")
    if not isinstance(activities, list) or not activities:
        raise SystemExit("Plan JSON must contain a non-empty 'activities' list.")
    names = []
    for index, activity in enumerate(activities):
        filename = activity.get("filename")
        if not isinstance(filename, str) or not filename.lower().endswith(".docx"):
            raise SystemExit(f"Plan activity {index} has no valid DOCX filename.")
        names.append(filename)
    return names


def validate_delivery(
    directory: Path,
    month: int | None,
    branch: str,
    expected_count: int | None,
    plan_path: Path | None,
) -> dict:
    files = find_docx(directory)
    errors = []
    warnings = []
    expected_names = load_expected_names(plan_path) if plan_path else []

    if expected_count is not None and len(files) != expected_count:
        errors.append(
            f"交付数量不符合预期: {len(files)} != {expected_count}。"
        )

    actual_names = {path.name for path in files}
    if expected_names:
        missing = sorted(set(expected_names) - actual_names)
        extra = sorted(actual_names - set(expected_names))
        if missing:
            errors.append(f"缺少计划文件: {missing}")
        if extra:
            errors.append(f"存在计划外 DOCX: {extra}")

    document_results = []
    Document = import_dependencies()
    for path in files:
        result = validate_docx(Document, path, month, branch)
        document_results.append(result)
        errors.extend(f"{path.name}: {message}" for message in result["errors"])
        warnings.extend(f"{path.name}: {message}" for message in result["warnings"])

    return {
        "directory": str(directory),
        "document_count": len(files),
        "expected_count": expected_count,
        "errors": errors,
        "warnings": warnings,
        "documents": document_results,
    }


def print_report(report: dict) -> None:
    print(f"交付目录: {report['directory']}")
    print(f"DOCX 数量: {report['document_count']}")
    for document in report["documents"]:
        states = []
        if document["errors"]:
            states.append(f"错误 {len(document['errors'])}")
        if document["warnings"]:
            states.append(f"警告 {len(document['warnings'])}")
        state_text = "，".join(states) if states else "通过"
        print(f"  {Path(document['path']).name}: {state_text}")
    for message in report["errors"]:
        print(f"错误: {message}")
    for message in report["warnings"]:
        print(f"警告: {message}")
    if not report["errors"]:
        print("结构验收通过。仍需逐页渲染并查看每一页。")


def main() -> None:
    configure_stdout()
    parser = argparse.ArgumentParser(
        description="Validate a monthly DOCX delivery folder."
    )
    parser.add_argument(
        "--directory",
        required=True,
        type=Path,
        help="Folder containing final DOCX files.",
    )
    parser.add_argument("--month", help="Target month such as '9月'.")
    parser.add_argument(
        "--branch-name",
        help="Party branch name expected in file names.",
    )
    parser.add_argument(
        "--expected-count",
        type=int,
        help="Expected number of DOCX files.",
    )
    parser.add_argument(
        "--plan",
        type=Path,
        help="Optional plan JSON produced by plan_monthly_docs.py.",
    )
    parser.add_argument(
        "--json",
        action="store_true",
        help="Print the validation report as JSON.",
    )
    args = parser.parse_args()

    directory = args.directory.resolve()
    if not directory.is_dir():
        raise SystemExit(f"Directory is not a folder: {directory}")
    if args.expected_count is not None and args.expected_count < 0:
        raise SystemExit("--expected-count must be zero or greater.")

    month = None
    if args.month:
        match = MONTH_RE.search(args.month)
        if not match:
            raise SystemExit(f"Month must contain a value such as '9月': {args.month}")
        month = int(match.group(1))
        if not 1 <= month <= 12:
            raise SystemExit(f"Month is outside 1 to 12: {args.month}")

    report = validate_delivery(
        directory,
        month,
        branch_label(args.branch_name),
        args.expected_count,
        args.plan.resolve() if args.plan else None,
    )
    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        print_report(report)
    if report["errors"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
