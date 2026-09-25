#!/usr/bin/env python
"""Inspect monthly party-building workbooks and DOCX record examples."""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path


RECORD_RE = re.compile(
    r"(党支部|支委会|党员大会|集中学习|主题党日|谈心谈话|讲党课)"
)


def configure_stdout() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")


def import_dependencies():
    try:
        import openpyxl
        from docx import Document
        from docx.oxml.ns import qn
    except ImportError as exc:
        raise SystemExit(
            "Missing dependency. Run this script with the bundled workspace "
            "Python that includes openpyxl and python-docx."
        ) from exc
    return openpyxl, Document, qn


def find_workbooks(root: Path, month: str | None) -> list[Path]:
    files = []
    for path in root.rglob("*.xlsx"):
        if path.name.startswith("~$"):
            continue
        text = str(path)
        if "重点工作清单" not in path.name:
            continue
        if month and month not in text:
            continue
        files.append(path)
    return sorted(files)


def find_docx_records(root: Path, month: str | None) -> list[Path]:
    files = []
    for path in root.rglob("*.docx"):
        if path.name.startswith("~$") or not RECORD_RE.search(path.name):
            continue
        if month and month not in str(path):
            continue
        files.append(path)
    return sorted(files)


def workbook_rows(openpyxl_module, path: Path) -> list[dict]:
    workbook = openpyxl_module.load_workbook(path, data_only=False)
    sheets = []
    for sheet in workbook.worksheets:
        rows = []
        for row_number in range(1, sheet.max_row + 1):
            values = [
                sheet.cell(row_number, column).value
                for column in range(1, sheet.max_column + 1)
            ]
            if any(value is not None for value in values):
                rows.append({"row": row_number, "values": values})
        sheets.append(
            {
                "name": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "rows": rows,
            }
        )
    return sheets


def paragraph_record(paragraph, qn_module) -> dict:
    paragraph_format = paragraph.paragraph_format
    first_run = paragraph.runs[0] if paragraph.runs else None
    east_asia_font = None
    if (
        first_run is not None
        and first_run._element.rPr is not None
        and first_run._element.rPr.rFonts is not None
    ):
        east_asia_font = first_run._element.rPr.rFonts.get(qn_module("w:eastAsia"))
    return {
        "style": paragraph.style.name,
        "alignment": str(paragraph_format.alignment),
        "first_line_indent": str(paragraph_format.first_line_indent),
        "line_spacing": str(paragraph_format.line_spacing),
        "font": first_run.font.name if first_run is not None else None,
        "east_asia_font": east_asia_font,
        "size_pt": (
            first_run.font.size.pt
            if first_run is not None and first_run.font.size is not None
            else None
        ),
        "bold": first_run.bold if first_run is not None else None,
        "text": paragraph.text,
    }


def docx_summary(Document, qn_module, path: Path, details: bool) -> dict:
    document = Document(path)
    paragraphs = [p for p in document.paragraphs if p.text.strip()]
    section = document.sections[0]
    summary = {
        "path": str(path),
        "name": path.name,
        "paragraphs": len(paragraphs),
        "characters": sum(len(p.text) for p in paragraphs),
        "tables": len(document.tables),
        "inline_shapes": len(document.inline_shapes),
        "page_width_emu": section.page_width,
        "page_height_emu": section.page_height,
        "margins_emu": {
            "top": section.top_margin,
            "bottom": section.bottom_margin,
            "left": section.left_margin,
            "right": section.right_margin,
        },
    }
    if details:
        summary["paragraph_details"] = [
            paragraph_record(paragraph, qn_module) for paragraph in paragraphs
        ]
    return summary


def print_text_report(workbooks, documents) -> None:
    for workbook in workbooks:
        print(f"\n===== 清单: {workbook['path']} =====")
        for sheet in workbook["sheets"]:
            print(
                f"[{sheet['name']}] "
                f"{sheet['max_row']} rows x {sheet['max_column']} columns"
            )
            for row in sheet["rows"]:
                print(f"ROW {row['row']}: {row['values']!r}")

    for document in documents:
        print(f"\n===== 记录: {document['name']} =====")
        print(
            f"paragraphs={document['paragraphs']} "
            f"characters={document['characters']} "
            f"tables={document['tables']} "
            f"inline_shapes={document['inline_shapes']}"
        )
        print(
            f"page={document['page_width_emu']}x{document['page_height_emu']} "
            f"margins={document['margins_emu']}"
        )
        for index, paragraph in enumerate(document.get("paragraph_details", [])):
            print(
                f"P{index:02d} style={paragraph['style']} "
                f"indent={paragraph['first_line_indent']} "
                f"font={paragraph['font']}/{paragraph['east_asia_font']} "
                f"size={paragraph['size_pt']} bold={paragraph['bold']} | "
                f"{paragraph['text']}"
            )


def main() -> None:
    configure_stdout()
    parser = argparse.ArgumentParser(
        description="Inspect monthly party-building checklists and DOCX records."
    )
    parser.add_argument(
        "--root",
        type=Path,
        default=Path.cwd(),
        help="Folder containing monthly checklists and record DOCX files.",
    )
    parser.add_argument(
        "--month",
        help="Optional filter such as '9月' or '2026年9月'.",
    )
    parser.add_argument(
        "--details",
        action="store_true",
        help="Include paragraph-level style and text details.",
    )
    parser.add_argument(
        "--json",
        action="store_true",
        help="Print JSON instead of a text report.",
    )
    args = parser.parse_args()

    root = args.root.resolve()
    if not root.is_dir():
        raise SystemExit(f"Root is not a directory: {root}")

    openpyxl, Document, qn = import_dependencies()
    workbooks = [
        {"path": str(path), "sheets": workbook_rows(openpyxl, path)}
        for path in find_workbooks(root, args.month)
    ]
    documents = [
        docx_summary(Document, qn, path, args.details)
        for path in find_docx_records(root, args.month)
    ]

    result = {
        "root": str(root),
        "month_filter": args.month,
        "workbooks": workbooks,
        "documents": documents,
    }
    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        print_text_report(workbooks, documents)


if __name__ == "__main__":
    main()
