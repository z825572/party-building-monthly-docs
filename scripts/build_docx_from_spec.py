#!/usr/bin/env python
"""Build DOCX records by copying local templates and replacing body paragraphs."""

from __future__ import annotations

import argparse
import json
import shutil
import sys
import uuid
import zipfile
from pathlib import Path


ALIGNMENTS = {
    "left": "LEFT",
    "center": "CENTER",
    "right": "RIGHT",
    "justify": "JUSTIFY",
}


def configure_stdout() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")


def import_dependencies():
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
    except ImportError as exc:
        raise SystemExit(
            "Missing dependency. Run this script with the bundled workspace "
            "Python that includes python-docx."
        ) from exc
    return Document, WD_ALIGN_PARAGRAPH, Pt


def load_spec(path: Path) -> dict:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise SystemExit(f"Cannot read JSON spec: {exc}") from exc
    if not isinstance(data, dict) or not isinstance(data.get("documents"), list):
        raise SystemExit("Spec must be an object with a non-empty 'documents' list.")
    if not data["documents"]:
        raise SystemExit("Spec contains no documents.")
    return data


def resolve_path(value: str, base_dir: Path) -> Path:
    path = Path(value)
    return path if path.is_absolute() else (base_dir / path).resolve()


def optional_number(value, prefix: str, field: str) -> float | None:
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise SystemExit(f"{prefix}.{field} must be a number or null.")
    if float(value) < 0:
        raise SystemExit(f"{prefix}.{field} must be zero or greater.")
    return float(value)


def optional_bool(
    value,
    prefix: str,
    field: str,
    default: bool | None = False,
) -> bool | None:
    if value is None:
        return default
    if not isinstance(value, bool):
        raise SystemExit(f"{prefix}.{field} must be true or false.")
    return value


def validate_document_spec(spec: dict, base_dir: Path, index: int) -> dict:
    prefix = f"documents[{index}]"
    if not isinstance(spec, dict):
        raise SystemExit(f"{prefix} must be an object.")
    for field in ("template", "filename", "paragraphs"):
        if field not in spec:
            raise SystemExit(f"{prefix} is missing '{field}'.")
    if not isinstance(spec["paragraphs"], list) or not spec["paragraphs"]:
        raise SystemExit(f"{prefix}.paragraphs must be a non-empty list.")

    template = resolve_path(str(spec["template"]), base_dir)
    if not template.is_file() or template.suffix.lower() != ".docx":
        raise SystemExit(f"{prefix}.template is not a DOCX file: {template}")

    filename = str(spec["filename"])
    if Path(filename).name != filename or not filename.lower().endswith(".docx"):
        raise SystemExit(f"{prefix}.filename must be a DOCX basename.")

    paragraphs = []
    for paragraph_index, paragraph in enumerate(spec["paragraphs"]):
        paragraph_prefix = f"{prefix}.paragraphs[{paragraph_index}]"
        if not isinstance(paragraph, dict):
            raise SystemExit(f"{paragraph_prefix} must be an object.")
        style = str(paragraph.get("style", "")).strip()
        if not style:
            raise SystemExit(f"{paragraph_prefix} is missing 'style'.")

        blank = optional_bool(paragraph.get("blank", False), paragraph_prefix, "blank")
        text = paragraph.get("text", "")
        if not isinstance(text, str):
            raise SystemExit(f"{paragraph_prefix}.text must be a string.")
        if not text.strip() and not blank:
            raise SystemExit(
                f"{paragraph_prefix} is empty; set 'blank': true for a blank paragraph."
            )
        if blank:
            text = ""
        alignment = paragraph.get("alignment")
        if alignment is not None and alignment not in ALIGNMENTS:
            raise SystemExit(
                f"{paragraph_prefix}.alignment must be one of "
                f"{sorted(ALIGNMENTS)}."
            )
        first_line_indent_pt = optional_number(
            paragraph.get("first_line_indent_pt"),
            paragraph_prefix,
            "first_line_indent_pt",
        )
        left_indent_pt = optional_number(
            paragraph.get("left_indent_pt"),
            paragraph_prefix,
            "left_indent_pt",
        )
        right_indent_pt = optional_number(
            paragraph.get("right_indent_pt"),
            paragraph_prefix,
            "right_indent_pt",
        )
        space_before_pt = optional_number(
            paragraph.get("space_before_pt"),
            paragraph_prefix,
            "space_before_pt",
        )
        space_after_pt = optional_number(
            paragraph.get("space_after_pt"),
            paragraph_prefix,
            "space_after_pt",
        )
        line_spacing_pt = optional_number(
            paragraph.get("line_spacing_pt"),
            paragraph_prefix,
            "line_spacing_pt",
        )
        paragraphs.append(
            {
                "style": style,
                "text": text,
                "blank": blank,
                "bold": optional_bool(
                    paragraph.get("bold"),
                    paragraph_prefix,
                    "bold",
                    default=None,
                ),
                "italic": optional_bool(
                    paragraph.get("italic"),
                    paragraph_prefix,
                    "italic",
                    default=None,
                ),
                "alignment": alignment,
                "first_line_indent_pt": first_line_indent_pt,
                "left_indent_pt": left_indent_pt,
                "right_indent_pt": right_indent_pt,
                "space_before_pt": space_before_pt,
                "space_after_pt": space_after_pt,
                "line_spacing_pt": line_spacing_pt,
                "page_break_before": optional_bool(
                    paragraph.get("page_break_before"),
                    paragraph_prefix,
                    "page_break_before",
                ),
                "keep_with_next": optional_bool(
                    paragraph.get("keep_with_next"),
                    paragraph_prefix,
                    "keep_with_next",
                ),
                "keep_together": optional_bool(
                    paragraph.get("keep_together"),
                    paragraph_prefix,
                    "keep_together",
                ),
            }
        )

    return {
        "template": template,
        "filename": filename,
        "title": str(spec.get("title", Path(filename).stem)),
        "paragraphs": paragraphs,
    }


def clear_paragraphs(document) -> None:
    for paragraph in list(document.paragraphs):
        paragraph._element.getparent().remove(paragraph._element)


def validate_output(Document, output_path: Path, paragraph_specs: list[dict]) -> None:
    try:
        with zipfile.ZipFile(output_path) as archive:
            bad_member = archive.testzip()
    except (OSError, zipfile.BadZipFile) as exc:
        raise SystemExit(f"Generated file is not a valid DOCX ZIP: {exc}") from exc
    if bad_member:
        raise SystemExit(f"Generated DOCX contains a corrupt ZIP member: {bad_member}")

    document = Document(output_path)
    if len(document.paragraphs) != len(paragraph_specs):
        raise SystemExit(
            "Generated DOCX paragraph count does not match the spec: "
            f"{len(document.paragraphs)} != {len(paragraph_specs)}"
        )
    for index, (paragraph, paragraph_spec) in enumerate(
        zip(document.paragraphs, paragraph_specs)
    ):
        if paragraph.text != paragraph_spec["text"]:
            raise SystemExit(
                f"Generated DOCX text mismatch at paragraph {index}."
            )


def build_document(
    Document,
    WD_ALIGN_PARAGRAPH,
    Pt,
    document_spec: dict,
    output_dir: Path,
    overwrite: bool,
    dry_run: bool,
) -> Path:
    output_path = output_dir / document_spec["filename"]
    if output_path.exists() and not overwrite:
        raise SystemExit(f"Output already exists; pass --overwrite: {output_path}")
    if output_path.resolve() == document_spec["template"].resolve():
        raise SystemExit("Output path must not overwrite the selected template.")
    if dry_run:
        return output_path

    output_dir.mkdir(parents=True, exist_ok=True)
    temporary_path = output_dir / (
        f".{output_path.stem}.{uuid.uuid4().hex}.tmp.docx"
    )
    try:
        shutil.copy2(document_spec["template"], temporary_path)
        document = Document(temporary_path)

        missing_styles = sorted(
            {
                paragraph["style"]
                for paragraph in document_spec["paragraphs"]
                if paragraph["style"] not in document.styles
            }
        )
        if missing_styles:
            raise SystemExit(
                f"Template {document_spec['template']} is missing styles: "
                f"{missing_styles}"
            )

        clear_paragraphs(document)
        for paragraph_spec in document_spec["paragraphs"]:
            paragraph = document.add_paragraph(style=paragraph_spec["style"])
            if not paragraph_spec["blank"]:
                run = paragraph.add_run(paragraph_spec["text"])
                if paragraph_spec["bold"] is not None:
                    run.bold = paragraph_spec["bold"]
                if paragraph_spec["italic"] is not None:
                    run.italic = paragraph_spec["italic"]

            paragraph_format = paragraph.paragraph_format
            alignment = paragraph_spec["alignment"]
            if alignment:
                paragraph_format.alignment = getattr(
                    WD_ALIGN_PARAGRAPH, ALIGNMENTS[alignment]
                )

            for field, property_name in (
                ("first_line_indent_pt", "first_line_indent"),
                ("left_indent_pt", "left_indent"),
                ("right_indent_pt", "right_indent"),
                ("space_before_pt", "space_before"),
                ("space_after_pt", "space_after"),
                ("line_spacing_pt", "line_spacing"),
            ):
                value = paragraph_spec[field]
                if value is not None:
                    setattr(paragraph_format, property_name, Pt(value))

            if (
                not paragraph_spec["blank"]
                and paragraph_spec["bold"] is True
                and paragraph_spec["first_line_indent_pt"] is None
            ):
                paragraph_format.first_line_indent = Pt(0)

            if paragraph_spec["page_break_before"]:
                paragraph_format.page_break_before = True
            if paragraph_spec["keep_with_next"]:
                paragraph_format.keep_with_next = True
            if paragraph_spec["keep_together"]:
                paragraph_format.keep_together = True

        document.core_properties.title = document_spec["title"]
        document.save(temporary_path)
        validate_output(Document, temporary_path, document_spec["paragraphs"])
        temporary_path.replace(output_path)
    except BaseException:
        temporary_path.unlink(missing_ok=True)
        raise

    return output_path


def main() -> None:
    configure_stdout()
    parser = argparse.ArgumentParser(
        description="Build DOCX records from a template-inheritance JSON spec."
    )
    parser.add_argument("--spec", required=True, type=Path)
    parser.add_argument(
        "--output-dir",
        type=Path,
        help="Override output_dir from the spec.",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Replace existing output files.",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Validate the spec without writing DOCX files.",
    )
    args = parser.parse_args()

    spec_path = args.spec.resolve()
    raw_spec = load_spec(spec_path)
    base_dir = spec_path.parent
    output_value = str(args.output_dir or raw_spec.get("output_dir", "output"))
    output_dir = resolve_path(output_value, base_dir)

    document_specs = [
        validate_document_spec(document, base_dir, index)
        for index, document in enumerate(raw_spec["documents"])
    ]
    filenames = [document_spec["filename"] for document_spec in document_specs]
    duplicate_filenames = sorted(
        {
            filename
            for filename in filenames
            if filenames.count(filename) > 1
        }
    )
    if duplicate_filenames:
        raise SystemExit(
            f"Spec contains duplicate output filenames: {duplicate_filenames}"
        )

    Document, WD_ALIGN_PARAGRAPH, Pt = import_dependencies()
    outputs = [
        build_document(
            Document,
            WD_ALIGN_PARAGRAPH,
            Pt,
            document_spec,
            output_dir,
            args.overwrite,
            args.dry_run,
        )
        for document_spec in document_specs
    ]

    action = "Validated" if args.dry_run else "Created"
    print(f"{action} {len(outputs)} document(s):")
    for output in outputs:
        print(output)


if __name__ == "__main__":
    main()
