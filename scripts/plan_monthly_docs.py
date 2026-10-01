#!/usr/bin/env python
"""Build a monthly DOCX delivery plan from a work checklist."""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

from inspect_monthly_inputs import (
    RECORD_RE,
    configure_stdout,
    find_workbooks,
    workbook_rows,
)


MONTH_RE = re.compile(r"(\d{1,2})月")
ORDINAL_RE = re.compile(r"第([一二三四五六七八九十两\d]+)次")
ORDINALS = {
    "一": 1,
    "二": 2,
    "两": 2,
    "三": 3,
    "四": 4,
    "五": 5,
    "六": 6,
    "七": 7,
    "八": 8,
    "九": 9,
    "十": 10,
}
PLACEHOLDER_RE = re.compile(r"x{2,}", re.IGNORECASE)


def import_openpyxl():
    try:
        import openpyxl
    except ImportError as exc:
        raise SystemExit(
            "Missing dependency. Run this script with the bundled workspace "
            "Python that includes openpyxl."
        ) from exc
    return openpyxl


def as_text(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def month_number(value: str) -> int:
    match = MONTH_RE.search(value)
    if not match:
        raise SystemExit(f"Month must contain a value such as '9月': {value}")
    number = int(match.group(1))
    if not 1 <= number <= 12:
        raise SystemExit(f"Month is outside 1 to 12: {value}")
    return number


def branch_label(value: str | None) -> str:
    label = as_text(value)
    if not label:
        return "xx党支部"
    if not label.endswith("党支部"):
        label += "党支部"
    return label


def ordinal_number(value: str) -> int | None:
    match = ORDINAL_RE.search(value)
    if not match:
        return None
    token = match.group(1)
    if token.isdigit():
        return int(token)
    if token in ORDINALS:
        return ORDINALS[token]
    return None


def category_terms(category: str) -> list[str]:
    terms = []
    if "谈心谈话" in category:
        terms.append("谈心谈话")
    if "集中学习" in category:
        terms.append("集中学习")
    if "支委会" in category or "党员大会" in category:
        terms.extend(["支部党员大会", "支委会", "党员大会"])
    if "主题党日" in category:
        terms.append("主题党日")
    if "讲党课" in category or "党课" in category:
        terms.extend(["讲党课", "党课"])
    return list(dict.fromkeys(terms))


def normalize_title(
    raw_title: str,
    category: str,
    branch: str,
    month: int,
) -> tuple[str, list[str]]:
    warnings = []
    title = as_text(raw_title)
    branch_stem = branch[:-3] if branch.endswith("党支部") else branch
    if not title:
        title = f"{branch}{month}月{category}"
        return title if title.lower().endswith(".docx") else title + ".docx", warnings

    title = re.sub(
        r"^x{1,}(?:支行)?党支部",
        branch,
        title,
        flags=re.IGNORECASE,
    )
    title = re.sub(
        rf"^{re.escape(branch_stem)}(?!党支部)",
        branch,
        title,
    )
    if branch not in title and branch_stem not in title:
        title = branch + title

    if not MONTH_RE.search(title):
        ordinal = ORDINAL_RE.search(title)
        if ordinal:
            title = title[: ordinal.start()] + f"{month}月" + title[ordinal.start() :]
        else:
            suffix = title[len(branch) :]
            title = f"{branch}{month}月{suffix}"

    if PLACEHOLDER_RE.search(title):
        warnings.append("标题仍有 xxxx 占位符，需要结合清单内容或用户要求确定主题。")
    if not title.lower().endswith(".docx"):
        title += ".docx"
    return title, warnings


def find_header(sheet) -> dict[str, int]:
    for row_number in range(1, sheet.max_row + 1):
        headers = {
            as_text(sheet.cell(row_number, column).value): column
            for column in range(1, sheet.max_column + 1)
        }
        if not any("活动类别" in header for header in headers):
            continue
        if not any("活动主题" in header for header in headers):
            continue
        result = {}
        for header, column in headers.items():
            if "活动类别" in header:
                result["category"] = column
            elif "活动主题" in header:
                result["title"] = column
            elif "活动内容" in header:
                result["content"] = column
            elif "参与人员" in header:
                result["participants"] = column
            elif "完成时间" in header:
                result["due"] = column
            elif "要求" in header:
                result["requirements"] = column
            elif "备注" in header:
                result["notes"] = column
        result["header_row"] = row_number
        return result
    raise SystemExit("Could not find the checklist header row.")


def first_number_before(sheet, row_number: int, category_column: int) -> str:
    for column in range(1, category_column):
        value = sheet.cell(row_number, column).value
        if value is not None and as_text(value):
            return as_text(value)
    return ""


def cell_text(sheet, row_number: int, column: int | None) -> str:
    if not column:
        return ""
    return as_text(sheet.cell(row_number, column).value)


def extract_activities(sheet) -> list[dict]:
    columns = find_header(sheet)
    mapped_columns = {
        column
        for key, column in columns.items()
        if key != "header_row" and isinstance(column, int)
    }
    activities = []
    sequence = 0
    for row_number in range(columns["header_row"] + 1, sheet.max_row + 1):
        category = as_text(sheet.cell(row_number, columns["category"]).value)
        if not category:
            continue
        content = cell_text(sheet, row_number, columns.get("content"))
        title = cell_text(sheet, row_number, columns.get("title"))
        if not content and not title:
            continue
        sequence += 1
        activities.append(
            {
                "sequence": sequence,
                "source_number": first_number_before(
                    sheet, row_number, columns["category"]
                ),
                "source_row": row_number,
                "category": category,
                "raw_title": title,
                "content": content,
                "participants": cell_text(
                    sheet, row_number, columns.get("participants")
                ),
                "due": cell_text(sheet, row_number, columns.get("due")),
                "requirements": cell_text(
                    sheet, row_number, columns.get("requirements")
                ),
                "notes": cell_text(sheet, row_number, columns.get("notes")),
                "extra_values": [
                    as_text(sheet.cell(row_number, column).value)
                    for column in range(1, sheet.max_column + 1)
                    if column not in mapped_columns
                    and as_text(sheet.cell(row_number, column).value)
                ],
            }
        )
    return activities


def candidate_month(path: Path) -> int | None:
    matches = MONTH_RE.findall(str(path))
    if not matches:
        return None
    return int(matches[-1])


def month_distance(target: int, candidate: int) -> int:
    return (target - candidate) % 12


def find_template_candidates(
    root: Path,
    category: str,
    title: str,
    target_month: int,
    lookback: int,
) -> list[dict]:
    terms = category_terms(category)
    required_terms = (
        ["谈心谈话", "集中学习"]
        if "谈心谈话" in category and "集中学习" in category
        else terms[:1]
    )
    title_ordinal = ordinal_number(title)
    candidates = []
    for path in root.rglob("*.docx"):
        if path.name.startswith("~$") or not RECORD_RE.search(path.name):
            continue
        month = candidate_month(path)
        if month is None:
            continue
        distance = month_distance(target_month, month)
        if distance < 1 or distance > lookback:
            continue

        score = 100 - distance * 20
        matched_terms = [term for term in terms if term in path.name]
        matched_required = [
            term for term in required_terms if term in path.name
        ]
        if required_terms:
            score += len(matched_required) * 50
            if len(matched_required) != len(required_terms):
                score -= 60
            score += (
                len([term for term in matched_terms if term not in required_terms])
                * 5
            )
        candidate_ordinal = ordinal_number(path.name)
        if title_ordinal is not None and candidate_ordinal == title_ordinal:
            score += 20
        if score <= 0:
            continue
        candidates.append(
            {
                "path": str(path),
                "name": path.name,
                "month": month,
                "month_distance": distance,
                "score": score,
                "matched_terms": matched_terms,
                "required_terms": required_terms,
            }
        )
    return sorted(
        candidates,
        key=lambda item: (-item["score"], item["month_distance"], item["name"]),
    )[:3]


def build_plan(root: Path, month: int, branch: str, lookback: int) -> dict:
    workbooks = find_workbooks(root, f"{month}月")
    if not workbooks:
        raise SystemExit(f"No checklist workbook found for {month}月 under {root}")
    workbook_path = workbooks[-1]
    workbook = import_openpyxl().load_workbook(workbook_path, data_only=True)
    sheet = workbook.worksheets[0]
    activities = extract_activities(sheet)
    if not activities:
        raise SystemExit(f"No activities found in {workbook_path}")

    filenames = defaultdict(list)
    for activity in activities:
        filename, warnings = normalize_title(
            activity["raw_title"],
            activity["category"],
            branch,
            month,
        )
        activity["filename"] = filename
        activity["warnings"] = warnings
        filenames[filename].append(activity["sequence"])
    for activity in activities:
        duplicates = filenames[activity["filename"]]
        if len(duplicates) > 1:
            activity["warnings"].append(
                "规范化后与其他活动重名，需要人工调整文件名。"
            )
        candidates = find_template_candidates(
            root,
            activity["category"],
            filename,
            month,
            lookback,
        )
        activity["template_candidates"] = candidates
        if not candidates:
            activity["warnings"].append(
                "未找到最近同类型模板，生成前需要人工指定模板。"
            )

    plan_warnings = []
    if branch == "xx党支部":
        plan_warnings.append(
            "未提供 --branch-name，文件名仍使用 xx党支部 占位符。"
        )
    return {
        "root": str(root),
        "workbook": str(workbook_path),
        "month": month,
        "branch": branch,
        "warnings": plan_warnings,
        "lookback_months": lookback,
        "activity_count": len(activities),
        "activities": activities,
    }


def print_text_plan(plan: dict) -> None:
    print(f"月份: {plan['month']}月")
    print(f"纲领文件: {plan['workbook']}")
    print(f"党支部: {plan['branch']}")
    print(f"活动数: {plan['activity_count']}")
    for warning in plan.get("warnings", []):
        print(f"警告: {warning}")
    for activity in plan["activities"]:
        print(
            f"\n[{activity['sequence']}] {activity['category']} "
            f"(清单行 {activity['source_row']})"
        )
        print(f"  输出: {activity['filename']}")
        print(f"  完成时间: {activity['due'] or '未填写'}")
        if activity["template_candidates"]:
            best = activity["template_candidates"][0]
            print(f"  推荐模板: {best['path']} (score={best['score']})")
            for candidate in activity["template_candidates"][1:]:
                print(
                    f"  备选模板: {candidate['path']} "
                    f"(score={candidate['score']})"
                )
        for warning in activity["warnings"]:
            print(f"  警告: {warning}")


def main() -> None:
    configure_stdout()
    parser = argparse.ArgumentParser(
        description="Build a monthly DOCX delivery plan from a checklist."
    )
    parser.add_argument(
        "--root",
        type=Path,
        default=Path.cwd(),
        help="Folder containing monthly checklists and historical DOCX records.",
    )
    parser.add_argument("--month", required=True, help="Target month such as '9月'.")
    parser.add_argument(
        "--branch-name",
        help="Party branch name. '党支部' is appended when omitted.",
    )
    parser.add_argument(
        "--lookback",
        type=int,
        default=3,
        help="Number of prior months to search for templates.",
    )
    parser.add_argument(
        "--json",
        action="store_true",
        help="Print the complete plan as JSON.",
    )
    args = parser.parse_args()

    root = args.root.resolve()
    if not root.is_dir():
        raise SystemExit(f"Root is not a directory: {root}")
    if args.lookback < 1 or args.lookback > 12:
        raise SystemExit("--lookback must be between 1 and 12.")

    plan = build_plan(
        root,
        month_number(args.month),
        branch_label(args.branch_name),
        args.lookback,
    )
    if args.json:
        print(json.dumps(plan, ensure_ascii=False, indent=2))
    else:
        print_text_plan(plan)


if __name__ == "__main__":
    main()
