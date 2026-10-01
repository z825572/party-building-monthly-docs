#!/usr/bin/env python
"""Profile page, paragraph, and run formatting in DOCX records."""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any


RECORD_RE = re.compile(
    r"(党支部|支委会|党员大会|集中学习|主题党日|谈心谈话|讲党课)"
)
MONTH_RE = re.compile(r"(\d{1,2})月")


def configure_stdout() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")


def import_dependencies():
    try:
        from docx import Document
        from docx.oxml.ns import qn
    except ImportError as exc:
        raise SystemExit(
            "Missing dependency. Run this script with the bundled workspace "
            "Python that includes python-docx."
        ) from exc
    return Document, qn


def find_docx(
    root: Path,
    months: list[str],
    include_source: bool,
    filename_pattern: re.Pattern[str] | None = None,
) -> list[Path]:
    files = []
    for path in root.rglob("*.docx"):
        if path.name.startswith("~$"):
            continue
        if months and not any(month in str(path) for month in months):
            continue
        if not include_source and not RECORD_RE.search(path.name):
            continue
        if filename_pattern and not filename_pattern.search(path.name):
            continue
        files.append(path)
    return sorted(files)


def integer(value: str | None) -> int | None:
    if value is None:
        return None
    try:
        return int(value)
    except ValueError:
        return None


def on_off_value(node, tag: str, qn) -> bool | None:
    if node is None:
        return None
    element = node.find(qn(tag))
    if element is None:
        return None
    value = element.get(qn("w:val"))
    if value is None:
        return True
    return value.lower() not in {"0", "false", "off", "no"}


def text_value(node, tag: str, qn, attribute: str = "w:val") -> str | None:
    if node is None:
        return None
    element = node.find(qn(tag))
    if element is None:
        return None
    return element.get(qn(attribute))


def read_paragraph_properties(properties, qn) -> dict[str, Any]:
    if properties is None:
        return {}
    result: dict[str, Any] = {}

    alignment = text_value(properties, "w:jc", qn)
    if alignment is not None:
        result["alignment"] = alignment

    spacing = properties.find(qn("w:spacing"))
    if spacing is not None:
        before = integer(spacing.get(qn("w:before")))
        after = integer(spacing.get(qn("w:after")))
        line = integer(spacing.get(qn("w:line")))
        rule = spacing.get(qn("w:lineRule")) or "auto"
        if before is not None:
            result["before_twips"] = before
            result["before_pt"] = round(before / 20, 2)
        if after is not None:
            result["after_twips"] = after
            result["after_pt"] = round(after / 20, 2)
        if line is not None:
            result["line_raw"] = line
            result["line_rule"] = rule
            if rule == "auto":
                result["line_multiple"] = round(line / 240, 3)
            else:
                result["line_pt"] = round(line / 20, 2)

    indent = properties.find(qn("w:ind"))
    if indent is not None:
        fields = {
            "firstLine": "first_line",
            "hanging": "hanging",
            "left": "left",
            "right": "right",
        }
        for xml_name, output_name in fields.items():
            value = integer(indent.get(qn(f"w:{xml_name}")))
            if value is not None:
                result[f"{output_name}_twips"] = value
                result[f"{output_name}_cm"] = round(value / 567, 3)

    for tag, name in (
        ("w:keepNext", "keep_next"),
        ("w:keepLines", "keep_lines"),
        ("w:pageBreakBefore", "page_break_before"),
        ("w:widowControl", "widow_control"),
        ("w:snapToGrid", "snap_to_grid"),
    ):
        value = on_off_value(properties, tag, qn)
        if value is not None:
            result[name] = value

    outline = properties.find(qn("w:outlineLvl"))
    if outline is not None:
        result["outline_level"] = integer(outline.get(qn("w:val")))
    return result


def read_run_properties(properties, qn) -> dict[str, Any]:
    if properties is None:
        return {}
    result: dict[str, Any] = {}

    fonts = properties.find(qn("w:rFonts"))
    if fonts is not None:
        for attribute, name in (
            ("w:ascii", "font_ascii"),
            ("w:eastAsia", "font_east_asia"),
            ("w:hAnsi", "font_hansi"),
            ("w:cs", "font_cs"),
            ("w:hint", "font_hint"),
        ):
            value = fonts.get(qn(attribute))
            if value is not None:
                result[name] = value

    size = integer(text_value(properties, "w:sz", qn))
    if size is not None:
        result["size_half_points"] = size
        result["size_pt"] = round(size / 2, 2)
    complex_size = integer(text_value(properties, "w:szCs", qn))
    if complex_size is not None:
        result["size_complex_half_points"] = complex_size
        result["size_complex_pt"] = round(complex_size / 2, 2)

    for tag, name in (
        ("w:b", "bold"),
        ("w:i", "italic"),
        ("w:strike", "strike"),
        ("w:smallCaps", "small_caps"),
        ("w:snapToGrid", "snap_to_grid"),
    ):
        value = on_off_value(properties, tag, qn)
        if value is not None:
            result[name] = value

    character_spacing = integer(text_value(properties, "w:spacing", qn))
    if character_spacing is not None:
        result["character_spacing_twips"] = character_spacing
        result["character_spacing_pt"] = round(character_spacing / 20, 2)

    kerning = integer(text_value(properties, "w:kern", qn))
    if kerning is not None:
        result["kerning_half_points"] = kerning
        result["kerning_pt"] = round(kerning / 2, 2)

    color = text_value(properties, "w:color", qn)
    if color is not None:
        result["color"] = color
    return result


def merged_properties(
    style_chain, direct_properties, qn, reader
) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for style in style_chain:
        style_properties = style.element.find(qn("w:pPr"))
        if reader is read_run_properties:
            style_properties = style.element.find(qn("w:rPr"))
        result.update(reader(style_properties, qn))
    result.update(reader(direct_properties, qn))
    return result


def style_chain(style) -> list:
    chain = []
    seen = set()
    current = style
    while current is not None:
        style_id = current.style_id
        if style_id in seen:
            break
        chain.append(current)
        seen.add(style_id)
        current = current.base_style
    return list(reversed(chain))


def style_definition(style, qn) -> dict[str, Any]:
    next_style = style.element.find(qn("w:next"))
    return {
        "style_id": style.style_id,
        "name": style.name,
        "type": str(style.type),
        "builtin": style.builtin,
        "base_style": style.base_style.name if style.base_style else None,
        "next_style": next_style.get(qn("w:val")) if next_style is not None else None,
        "paragraph": read_paragraph_properties(
            style.element.find(qn("w:pPr")), qn
        ),
        "run": read_run_properties(style.element.find(qn("w:rPr")), qn),
    }


def section_layout(section, qn) -> dict[str, Any]:
    sect_pr = section._sectPr
    grid = sect_pr.find(qn("w:docGrid"))
    return {
        "page_width_cm": round(section.page_width.cm, 2),
        "page_height_cm": round(section.page_height.cm, 2),
        "orientation": str(section.orientation),
        "top_margin_cm": round(section.top_margin.cm, 2),
        "bottom_margin_cm": round(section.bottom_margin.cm, 2),
        "left_margin_cm": round(section.left_margin.cm, 2),
        "right_margin_cm": round(section.right_margin.cm, 2),
        "header_distance_cm": round(section.header_distance.cm, 2),
        "footer_distance_cm": round(section.footer_distance.cm, 2),
        "document_grid_type": (
            grid.get(qn("w:type")) if grid is not None else None
        ),
        "document_grid_line_pitch": (
            integer(grid.get(qn("w:linePitch"))) if grid is not None else None
        ),
    }


def compact_paragraph(paragraph, qn) -> dict[str, Any]:
    chain = style_chain(paragraph.style)
    paragraph_format = merged_properties(
        chain, paragraph._p.pPr, qn, read_paragraph_properties
    )
    base_run = merged_properties(chain, None, qn, read_run_properties)

    run_formats = []
    seen = set()
    for run in paragraph.runs:
        run_format = dict(base_run)
        run_format.update(read_run_properties(run._element.rPr, qn))
        signature = json.dumps(run_format, ensure_ascii=False, sort_keys=True)
        if signature not in seen:
            seen.add(signature)
            run_formats.append(run_format)

    return {
        "style": paragraph.style.name,
        "characters": len(paragraph.text),
        "paragraph_format": paragraph_format,
        "run_formats": run_formats,
    }


def document_profile(Document, qn, path: Path, root: Path) -> dict[str, Any]:
    document = Document(path)
    kind = "record" if RECORD_RE.search(path.name) else "source"
    all_paragraphs = document.paragraphs
    nonempty = [paragraph for paragraph in all_paragraphs if paragraph.text.strip()]
    empty = [paragraph for paragraph in all_paragraphs if not paragraph.text.strip()]

    style_counts = Counter(paragraph.style.name for paragraph in nonempty)
    paragraph_formats: dict[str, Counter] = defaultdict(Counter)
    run_formats: dict[str, Counter] = defaultdict(Counter)
    style_definitions: dict[str, dict[str, Any]] = {}

    for paragraph in nonempty:
        style_name = paragraph.style.name
        details = compact_paragraph(paragraph, qn)
        paragraph_formats[style_name][
            json.dumps(details["paragraph_format"], ensure_ascii=False, sort_keys=True)
        ] += 1
        for run_format in details["run_formats"]:
            run_formats[style_name][
                json.dumps(run_format, ensure_ascii=False, sort_keys=True)
            ] += 1
        for style in style_chain(paragraph.style):
            style_definitions[style.style_id] = style_definition(style, qn)

    month_match = MONTH_RE.search(str(path.relative_to(root)))
    return {
        "path": str(path.relative_to(root)),
        "name": path.name,
        "month": int(month_match.group(1)) if month_match else None,
        "kind": kind,
        "all_paragraph_count": len(all_paragraphs),
        "paragraph_count": len(nonempty),
        "empty_paragraph_count": len(empty),
        "empty_style_counts": dict(
            Counter(paragraph.style.name for paragraph in empty)
        ),
        "text_characters": sum(len(paragraph.text) for paragraph in nonempty),
        "table_count": len(document.tables),
        "inline_shape_count": len(document.inline_shapes),
        "sections": [
            section_layout(section, qn) for section in document.sections
        ],
        "style_counts": dict(style_counts),
        "paragraph_formats": {
            style_name: [
                {"count": count, "format": json.loads(value)}
                for value, count in counter.most_common()
            ]
            for style_name, counter in paragraph_formats.items()
        },
        "run_formats": {
            style_name: [
                {"count": count, "format": json.loads(value)}
                for value, count in counter.most_common()
            ]
            for style_name, counter in run_formats.items()
        },
        "style_definitions": style_definitions,
    }


def aggregate_profiles(profiles: list[dict[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {
        "document_counts": Counter(),
        "style_counts": defaultdict(Counter),
        "paragraph_formats": defaultdict(lambda: defaultdict(Counter)),
        "run_formats": defaultdict(lambda: defaultdict(Counter)),
        "sections": defaultdict(Counter),
    }
    for profile in profiles:
        kind = profile["kind"]
        result["document_counts"][kind] += 1
        for style_name, count in profile["style_counts"].items():
            result["style_counts"][kind][style_name] += count
        for style_name, formats in profile["paragraph_formats"].items():
            for entry in formats:
                signature = json.dumps(
                    entry["format"], ensure_ascii=False, sort_keys=True
                )
                result["paragraph_formats"][kind][style_name][signature] += entry[
                    "count"
                ]
        for style_name, formats in profile["run_formats"].items():
            for entry in formats:
                signature = json.dumps(
                    entry["format"], ensure_ascii=False, sort_keys=True
                )
                result["run_formats"][kind][style_name][signature] += entry["count"]
        for section in profile["sections"]:
            signature = json.dumps(section, ensure_ascii=False, sort_keys=True)
            result["sections"][kind][signature] += 1

    serializable = {
        "document_counts": dict(result["document_counts"]),
        "style_counts": {
            kind: dict(counter)
            for kind, counter in result["style_counts"].items()
        },
        "paragraph_formats": {
            kind: {
                style_name: [
                    {"count": count, "format": json.loads(value)}
                    for value, count in counter.most_common()
                ]
                for style_name, counter in styles.items()
            }
            for kind, styles in result["paragraph_formats"].items()
        },
        "run_formats": {
            kind: {
                style_name: [
                    {"count": count, "format": json.loads(value)}
                    for value, count in counter.most_common()
                ]
                for style_name, counter in styles.items()
            }
            for kind, styles in result["run_formats"].items()
        },
        "sections": {
            kind: [
                {"count": count, "layout": json.loads(value)}
                for value, count in counter.most_common()
            ]
            for kind, counter in result["sections"].items()
        },
    }
    return serializable


def print_report(profiles: list[dict[str, Any]]) -> None:
    for profile in profiles:
        print(
            f"{profile['month']}月 {profile['kind']} | {profile['name']} | "
            f"paragraphs={profile['paragraph_count']} "
            f"characters={profile['text_characters']} "
            f"tables={profile['table_count']} "
            f"shapes={profile['inline_shape_count']}"
        )
        print(f"  styles: {profile['style_counts']}")
        for section in profile["sections"]:
            print(f"  section: {section}")

    aggregate = aggregate_profiles(profiles)
    print("\n===== AGGREGATE =====")
    print(json.dumps(aggregate, ensure_ascii=False, indent=2))


def main() -> None:
    configure_stdout()
    parser = argparse.ArgumentParser(
        description="Profile page, style, paragraph, and run formatting in DOCX files."
    )
    parser.add_argument(
        "--root",
        type=Path,
        default=Path.cwd(),
        help="Folder containing monthly DOCX files.",
    )
    parser.add_argument(
        "--month",
        action="append",
        default=[],
        help="Month filter such as '6月'. Repeat to include multiple months.",
    )
    parser.add_argument(
        "--include-source",
        action="store_true",
        help="Also include DOCX files that do not look like formal records.",
    )
    parser.add_argument(
        "--filename-regex",
        help="Optional regular expression matched against the DOCX filename.",
    )
    parser.add_argument(
        "--json",
        action="store_true",
        help="Print per-file profiles and aggregate data as JSON.",
    )
    args = parser.parse_args()

    root = args.root.resolve()
    if not root.is_dir():
        raise SystemExit(f"Root is not a directory: {root}")

    Document, qn = import_dependencies()
    try:
        filename_pattern = (
            re.compile(args.filename_regex) if args.filename_regex else None
        )
    except re.error as exc:
        raise SystemExit(f"Invalid --filename-regex: {exc}") from exc
    profiles = [
        document_profile(Document, qn, path, root)
        for path in find_docx(
            root, args.month, args.include_source, filename_pattern
        )
    ]
    if args.json:
        print(
            json.dumps(
                {
                    "root": str(root),
                    "months": args.month,
                    "documents": profiles,
                    "aggregate": aggregate_profiles(profiles),
                },
                ensure_ascii=False,
                indent=2,
            )
        )
    else:
        print_report(profiles)


if __name__ == "__main__":
    main()
